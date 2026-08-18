"""College synopsis PDF + DOCX — real stack only, zero MERN mentions."""
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
)
from reportlab.lib import colors

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PDF_OUT = r"c:\Users\VishwajeetUsnale\Downloads\Project_Synopsis_IT_Corrected.pdf"
DOCX_OUT = r"c:\Users\VishwajeetUsnale\Downloads\Project_Synopsis_IT_Corrected.docx"
# Also overwrite / replace next to their original synopsis name
DOCX_OUT2 = r"c:\Users\VishwajeetUsnale\Downloads\Project Synopsis Corrected.docx"

TITLE = "AI-Powered SQL Agent & Gmail Auto-Reply Bot"
SUBTITLE = "Using React.js, Node.js, Express.js, MySQL & Google Gemini"
STACK = "React.js, Node.js, Express.js and MySQL"

PAGE_W, PAGE_H = A4
MARGIN = 0.9 * inch


def make_styles():
    base = getSampleStyleSheet()
    return {
        "college": ParagraphStyle(
            "college", parent=base["Normal"], fontName="Times-Bold",
            fontSize=13, alignment=TA_CENTER, leading=16, spaceAfter=4
        ),
        "affil": ParagraphStyle(
            "affil", parent=base["Normal"], fontName="Times-Roman",
            fontSize=11, alignment=TA_CENTER, leading=14, spaceAfter=2
        ),
        "label": ParagraphStyle(
            "label", parent=base["Normal"], fontName="Times-Bold",
            fontSize=12, alignment=TA_CENTER, spaceBefore=14, spaceAfter=6
        ),
        "title": ParagraphStyle(
            "title", parent=base["Normal"], fontName="Times-Bold",
            fontSize=16, alignment=TA_CENTER, leading=20, spaceBefore=6, spaceAfter=4
        ),
        "sub": ParagraphStyle(
            "sub", parent=base["Normal"], fontName="Times-Italic",
            fontSize=11, alignment=TA_CENTER, leading=14, spaceAfter=10
        ),
        "body": ParagraphStyle(
            "body", parent=base["Normal"], fontName="Times-Roman",
            fontSize=11, alignment=TA_JUSTIFY, leading=16, spaceAfter=8
        ),
        "h1": ParagraphStyle(
            "h1", parent=base["Normal"], fontName="Times-Bold",
            fontSize=14, alignment=TA_CENTER, spaceBefore=4, spaceAfter=12
        ),
        "h2": ParagraphStyle(
            "h2", parent=base["Normal"], fontName="Times-Bold",
            fontSize=12, alignment=TA_LEFT, spaceBefore=8, spaceAfter=6
        ),
        "bullet": ParagraphStyle(
            "bullet", parent=base["Normal"], fontName="Times-Roman",
            fontSize=11, alignment=TA_LEFT, leading=15, leftIndent=18, spaceAfter=4
        ),
        "center": ParagraphStyle(
            "center", parent=base["Normal"], fontName="Times-Roman",
            fontSize=11, alignment=TA_CENTER, leading=15, spaceAfter=4
        ),
        "mono": ParagraphStyle(
            "mono", parent=base["Normal"], fontName="Courier",
            fontSize=8.5, alignment=TA_LEFT, leading=11, spaceAfter=2
        ),
    }


def add_page_number(canvas, doc):
    canvas.saveState()
    canvas.setFont("Times-Roman", 9)
    canvas.drawCentredString(PAGE_W / 2, 0.55 * inch, str(canvas.getPageNumber()))
    canvas.restoreState()


def build_pdf():
    styles = make_styles()
    doc = SimpleDocTemplate(
        PDF_OUT, pagesize=A4,
        leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=0.75 * inch, bottomMargin=0.75 * inch,
        title=TITLE,
        author="Usnale Monika Prashant & Patil Vedanti Shahaji",
    )
    story = []

    # COVER
    story.append(Paragraph("Royal Education Society’s", styles["college"]))
    story.append(Paragraph(
        "College of Computer Science and Information Technology, Latur.",
        styles["college"]
    ))
    story.append(Paragraph("Affiliated to", styles["affil"]))
    story.append(Paragraph(
        "Swami Ramanand Teerth Marathwada University, Nanded.",
        styles["affil"]
    ))
    story.append(Spacer(1, 28))
    story.append(Paragraph("A Project Synopsis", styles["label"]))
    story.append(Paragraph("On", styles["center"]))
    story.append(Paragraph(TITLE, styles["title"]))
    story.append(Paragraph(SUBTITLE, styles["sub"]))
    story.append(Spacer(1, 16))
    story.append(Paragraph("Submitted for the award of degree of", styles["center"]))
    story.append(Paragraph("<b>Bachelor of Science (Information Technology)</b>", styles["center"]))
    story.append(Spacer(1, 18))
    story.append(Paragraph("<b>By:</b>", styles["center"]))
    story.append(Paragraph("Usnale Monika Prashant: 59(A)", styles["center"]))
    story.append(Paragraph("Patil Vedanti Shahaji: 38 (A)", styles["center"]))
    story.append(Spacer(1, 18))
    story.append(Paragraph("<b>Year: 2026 – 2027</b>", styles["center"]))
    story.append(Spacer(1, 40))
    sig = Table(
        [[
            Paragraph("_______________________<br/><b>Project Guide</b><br/>(Dr. Dipali Hemant Mahamuni)", styles["center"]),
            Paragraph("_______________________<br/><b>HOD</b>", styles["center"]),
        ]],
        colWidths=[3.2 * inch, 3.2 * inch]
    )
    story.append(sig)
    story.append(PageBreak())

    # TOC
    story.append(Paragraph("Table of Contents", styles["h1"]))
    toc_data = [
        ["Sr. No.", "Contents", "Page No."],
        ["1", "Abstract", "3"],
        ["2", "Introduction of Project", "4"],
        ["3", "Project Modules", "5"],
        ["4", "Project Plan (Gantt Chart)", "6"],
        ["5", "Project Requirements", "7"],
        ["6", "E-R Diagram", "8"],
        ["7", "Data Flow Diagram (DFD)", "9"],
        ["8", "Conclusion", "10"],
        ["9", "Bibliography", "11"],
    ]
    toc = Table(toc_data, colWidths=[0.9 * inch, 4.2 * inch, 1.0 * inch])
    toc.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"),
        ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E8E8")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (2, 0), (2, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(toc)
    story.append(PageBreak())

    # ABSTRACT
    story.append(Paragraph("1. Abstract", styles["h1"]))
    story.append(Paragraph(
        "Artificial Intelligence (AI) and Large Language Models (LLMs) are changing the way "
        "people use web applications and manage daily tasks. Many users find it difficult to "
        "write SQL queries or reply to emails manually. This project, <b>" + TITLE + "</b>, "
        "solves these problems by providing an easy-to-use system.",
        styles["body"]
    ))
    story.append(Paragraph(
        "The project is a full-stack web application developed using <b>" + STACK + "</b>. "
        "It includes an AI-powered SQL Agent that allows users to interact with the MySQL "
        "database using simple English. Users can add, update, delete, or view records without "
        "writing SQL queries. The system converts the user’s request into SQL commands, "
        "performs the required database operations, and generates PDF reports.",
        styles["body"]
    ))
    story.append(Paragraph(
        "The project also includes a Telegram Bot and Gmail Assistant. The bot checks unread "
        "Gmail messages, understands the email content using Google Gemini AI, and creates a "
        "suitable reply. Before sending the email, users can review, edit, approve, or reject "
        "the reply through Telegram, improving accuracy and security.",
        styles["body"]
    ))
    story.append(Paragraph(
        "Overall, this project combines AI, database management, email automation, and web "
        "technologies into a single application, helping users manage databases and emails "
        "more efficiently while reducing manual work.",
        styles["body"]
    ))
    story.append(PageBreak())

    # INTRODUCTION
    story.append(Paragraph("2. Introduction of Project", styles["h1"]))
    story.append(Paragraph(
        "Artificial Intelligence (AI) and Large Language Models (LLMs) are changing the way "
        "people use web applications. Many users find it difficult to write SQL queries and "
        "reply to emails manually. These tasks take time and require technical knowledge.",
        styles["body"]
    ))
    story.append(Paragraph(
        "The <b>" + TITLE + "</b> is developed to solve these problems. It allows users to "
        "interact with a MySQL database using simple English instead of writing SQL queries. "
        "The system understands the user’s request, converts it into an SQL query, and shows "
        "the required result.",
        styles["body"]
    ))
    story.append(Paragraph(
        "The project also includes a Gmail Auto-Reply Bot connected with Telegram. It checks "
        "unread emails, creates AI-generated reply drafts, and lets the user review and approve "
        "the reply before sending it. This helps save time while keeping the user in control.",
        styles["body"]
    ))
    story.append(Paragraph(
        "This project is developed using <b>" + STACK + "</b>, Google Gemini AI, Telegram Bot "
        "API, and Gmail API. It provides a simple, fast, and user-friendly solution for "
        "database management and email automation.",
        styles["body"]
    ))
    story.append(Paragraph("<b>Technology Used:</b>", styles["h2"]))
    for item in [
        "• Frontend: React.js (Vite), HTML, CSS, JavaScript",
        "• Backend: Node.js and Express.js",
        "• Database: MySQL",
        "• AI Technology: Google Gemini AI",
        "• APIs: Gmail API, Telegram Bot API",
    ]:
        story.append(Paragraph(item, styles["bullet"]))
    story.append(PageBreak())

    # MODULES
    story.append(Paragraph("3. Project Modules", styles["h1"]))
    modules = [
        ("1. User Authentication Module",
         "This module allows users to register and log in securely. Only authorized users can "
         "access the system and use its features."),
        ("2. AI SQL Agent Module",
         "This module allows users to interact with the database using simple English. The AI "
         "converts the user’s request into an SQL query, executes it, and displays the result."),
        ("3. Database Management Module",
         "This module manages all database operations such as adding, updating, deleting, and "
         "viewing records in the MySQL database."),
        ("4. Gmail Auto-Reply Module",
         "This module checks unread Gmail messages and generates reply drafts using Google Gemini AI."),
        ("5. Telegram Bot Module",
         "This module sends unread email details and AI-generated replies to Telegram. The user "
         "can review the reply and choose to send or skip it."),
        ("6. PDF Report Module",
         "This module generates PDF reports of the database records, allowing users to download "
         "and save the data easily."),
    ]
    for head, text in modules:
        story.append(Paragraph(head, styles["h2"]))
        story.append(Paragraph(text, styles["body"]))
    story.append(PageBreak())

    # GANTT
    story.append(Paragraph("4. Project Plan (Gantt Chart)", styles["h1"]))
    story.append(Paragraph(
        "The project follows a structured development plan across major phases:",
        styles["body"]
    ))
    gantt = [
        ["Sr.", "Task Name", "W1-2", "W3-4", "W5-6", "W7-8", "W9-10", "W11-12"],
        ["1", "Requirement Gathering", "■■■■", "", "", "", "", ""],
        ["2", "Planning", "■■", "■■", "", "", "", ""],
        ["3", "Designing", "", "■■■■", "■■", "", "", ""],
        ["4", "Coding / Development", "", "", "■■■■", "■■■■", "■■■■", ""],
        ["5", "Testing & Deployment", "", "", "", "", "■■", "■■■■"],
    ]
    gt = Table(gantt, colWidths=[0.45*inch, 1.9*inch, 0.6*inch, 0.6*inch, 0.6*inch, 0.6*inch, 0.65*inch, 0.7*inch])
    gt.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Times-Roman"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E8E8")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (2, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(gt)
    story.append(Spacer(1, 12))
    story.append(Paragraph("W1–W12 = Week 1 to Week 12 of the project timeline.", styles["center"]))
    story.append(PageBreak())

    # REQUIREMENTS
    story.append(Paragraph("5. Project Requirements", styles["h1"]))
    story.append(Paragraph("Hardware Requirements:", styles["h2"]))
    for item in [
        "1. Processor: Intel Core i5 (or equivalent)",
        "2. RAM: 8 GB minimum (16 GB recommended)",
        "3. Storage: 256 GB SSD (or higher)",
        "4. Input Device: Keyboard, Mouse",
        "5. Output Device: Monitor / LCD / LED, Printer",
        "6. Internet Connection: Required for Gemini AI, Gmail API, and Telegram Bot",
    ]:
        story.append(Paragraph(item, styles["bullet"]))

    story.append(Paragraph("Software Requirements:", styles["h2"]))
    for item in [
        "• Frontend: React.js, HTML, CSS, JavaScript, Vite, Visual Studio Code",
        "• Programming Language: JavaScript (Node.js)",
        "• Framework: Express.js, React.js",
        "• Database: MySQL",
        "• AI Technology: Google Gemini AI",
        "• API: Gmail API, Telegram Bot API",
        "• Platform: Web-Based Application",
    ]:
        story.append(Paragraph(item, styles["bullet"]))
    story.append(PageBreak())

    # ER
    story.append(Paragraph("6. E-R Diagram", styles["h1"]))
    story.append(Paragraph(
        "The system uses relational entities centered on the authenticated USER. "
        "Database records are managed through the AI SQL Agent. Email and AI reply entities "
        "support the Gmail–Telegram approval flow.",
        styles["body"]
    ))
    er = """
+-------------------+
|       USER        |
+-------------------+
| User_ID (PK)      |
| Name              |
| Email             |
| Password (hash)   |
+-------------------+
         |
         | 1:M
         v
+-------------------+          +-------------------+
|   DB_RECORD       |          |      EMAIL        |
| (e.g. students)   |          +-------------------+
+-------------------+          | Email_ID (PK)     |
| Record_ID (PK)    |          | User_ID (FK)      |
| User_ID (FK)      |          | From / Subject    |
| Name, Course...   |          +-------------------+
+-------------------+                    |
         |                               | 1:1
         | 1:M                           v
         v                    +-------------------+
+-------------------+         |    AI_REPLY       |
|   PDF_REPORT      |         +-------------------+
+-------------------+         | Reply_ID (PK)     |
| Report_ID (PK)    |         | Email_ID (FK)     |
| Record_ID (FK)    |         | Draft Text        |
+-------------------+         +-------------------+
                                       |
                                       | 1:M
                                       v
                              +-------------------+
                              |  TELEGRAM_ACTION  |
                              +-------------------+
                              | Telegram_ID (PK)  |
                              | Reply_ID (FK)     |
                              | Approve / Reject  |
                              +-------------------+
"""
    for line in er.strip("\n").split("\n"):
        story.append(Paragraph(line.replace(" ", "&nbsp;"), styles["mono"]))
    story.append(PageBreak())

    # DFD
    story.append(Paragraph("7. Data Flow Diagram (DFD)", styles["h1"]))
    story.append(Paragraph(
        "High-level data flow of the AI SQL Agent and Gmail–Telegram assistant:",
        styles["body"]
    ))
    dfd = """
                    +--------+
                    |  USER  |
                    +--------+
                         |
          Natural language query / Email review
                         |
                         v
     +----------------------------------------------+
     |  AI SQL Agent & Gmail Auto-Reply System      |
     |  (React Frontend + Node/Express Backend)     |
     +----------------------------------------------+
        |            |            |            |
        v            v            v            v
   +----------+ +---------+ +---------+ +------------+
   |  MySQL   | | Gemini  | |  Gmail  | | PDF Report |
   | Database | |   AI    | |   API   | |  Module    |
   +----------+ +---------+ +---------+ +------------+
                         |
              AI-generated email reply
                         |
                         v
                 +---------------+
                 | Telegram Bot  |
                 +---------------+
                         |
              Approve / Reject
                         |
                         v
                      USER
"""
    for line in dfd.strip("\n").split("\n"):
        story.append(Paragraph(line.replace(" ", "&nbsp;"), styles["mono"]))
    story.append(PageBreak())

    # CONCLUSION
    story.append(Paragraph("8. Conclusion", styles["h1"]))
    story.append(Paragraph(
        "The <b>" + TITLE + "</b> is a smart and user-friendly system that makes database "
        "management and email handling easier. It allows users to interact with a MySQL "
        "database using simple English without writing SQL queries. The system also generates "
        "AI-powered email replies and lets users review them on Telegram before sending.",
        styles["body"]
    ))
    story.append(Paragraph(
        "This project saves time, reduces manual work, and improves accuracy. It combines "
        "<b>" + STACK + "</b>, Google Gemini AI, Gmail API, and Telegram Bot API to provide "
        "an efficient and secure solution. The project can be further improved by adding more "
        "AI features and support for other messaging platforms in the future.",
        styles["body"]
    ))
    story.append(PageBreak())

    # BIBLIOGRAPHY
    story.append(Paragraph("9. Bibliography", styles["h1"]))
    story.append(Paragraph("<b>Books</b>", styles["h2"]))
    story.append(Paragraph(
        "1. Book: Systems Analysis and Design Methods<br/>"
        "Author: Jeffrey L. Whitten &amp; Lonnie D. Bentley<br/>"
        "Publication: McGraw-Hill Education",
        styles["body"]
    ))
    story.append(Paragraph(
        "2. Book: Learning React<br/>"
        "Author: Alex Banks &amp; Eve Porcello<br/>"
        "Publication: O’Reilly Media",
        styles["body"]
    ))
    story.append(Paragraph(
        "3. Book: MySQL Developer’s Library<br/>"
        "Author: Paul DuBois<br/>"
        "Publication: Addison-Wesley Professional",
        styles["body"]
    ))
    story.append(Paragraph("<b>Websites</b>", styles["h2"]))
    for w in [
        "React Documentation – https://react.dev",
        "Node.js Documentation – https://nodejs.org",
        "Express.js Documentation – https://expressjs.com",
        "Google Gemini AI Documentation – https://ai.google.dev",
        "MySQL Documentation – https://dev.mysql.com",
        "Gmail API Documentation – https://developers.google.com/gmail/api",
        "Telegram Bot API Documentation – https://core.telegram.org/bots/api",
    ]:
        story.append(Paragraph("• " + w, styles["bullet"]))

    doc.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)
    print("PDF:", PDF_OUT)


def add_center(doc, text, bold=False, size=12, space_after=6, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_body(doc, text, size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_after = Pt(8)
    return p


def add_heading_center(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = "Times New Roman"
    return p


def build_docx():
    d = Document()
    section = d.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # Cover
    add_center(d, "Royal Education Society’s", bold=True, size=13)
    add_center(d, "College of Computer Science and Information Technology, Latur.", bold=True, size=13)
    add_center(d, "Affiliated to", size=11)
    add_center(d, "Swami Ramanand Teerth Marathwada University, Nanded.", size=11)
    d.add_paragraph()
    add_center(d, "A Project Synopsis", bold=True, size=12)
    add_center(d, "On", size=11)
    add_center(d, TITLE, bold=True, size=16)
    add_center(d, SUBTITLE, size=11, italic=True)
    d.add_paragraph()
    add_center(d, "Submitted for the award of degree of", size=11)
    add_center(d, "Bachelor of Science (Information Technology)", bold=True, size=12)
    d.add_paragraph()
    add_center(d, "By:", bold=True, size=11)
    add_center(d, "Usnale Monika Prashant: 59(A)", size=11)
    add_center(d, "Patil Vedanti Shahaji: 38 (A)", size=11)
    d.add_paragraph()
    add_center(d, "Year: 2026 – 2027", bold=True, size=11)
    d.add_paragraph()
    d.add_paragraph()
    add_center(d, "_______________________                    _______________________", size=11)
    add_center(d, "Project Guide                                          HOD", bold=True, size=11)
    add_center(d, "(Dr. Dipali Hemant Mahamuni)", size=10)
    d.add_page_break()

    # TOC
    add_heading_center(d, "Table of Contents")
    toc = [
        ("1", "Abstract", "3"),
        ("2", "Introduction of Project", "4"),
        ("3", "Project Modules", "5"),
        ("4", "Project Plan (Gantt Chart)", "6"),
        ("5", "Project Requirements", "7"),
        ("6", "E-R Diagram", "8"),
        ("7", "Data Flow Diagram (DFD)", "9"),
        ("8", "Conclusion", "10"),
        ("9", "Bibliography", "11"),
    ]
    table = d.add_table(rows=1 + len(toc), cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Sr. No.", "Contents", "Page No."
    for i, (a, b, c) in enumerate(toc, start=1):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b
        table.rows[i].cells[2].text = c
    d.add_page_break()

    # Abstract
    add_heading_center(d, "1. Abstract")
    add_body(d,
        "Artificial Intelligence (AI) and Large Language Models (LLMs) are changing the way "
        "people use web applications and manage daily tasks. Many users find it difficult to "
        f"write SQL queries or reply to emails manually. This project, {TITLE}, solves these "
        "problems by providing an easy-to-use system."
    )
    add_body(d,
        f"The project is a full-stack web application developed using {STACK}. "
        "It includes an AI-powered SQL Agent that allows users to interact with the MySQL "
        "database using simple English. Users can add, update, delete, or view records without "
        "writing SQL queries. The system converts the user’s request into SQL commands, "
        "performs the required database operations, and generates PDF reports."
    )
    add_body(d,
        "The project also includes a Telegram Bot and Gmail Assistant. The bot checks unread "
        "Gmail messages, understands the email content using Google Gemini AI, and creates a "
        "suitable reply. Before sending the email, users can review, edit, approve, or reject "
        "the reply through Telegram, improving accuracy and security."
    )
    add_body(d,
        "Overall, this project combines AI, database management, email automation, and web "
        "technologies into a single application, helping users manage databases and emails "
        "more efficiently while reducing manual work."
    )
    d.add_page_break()

    # Introduction
    add_heading_center(d, "2. Introduction of Project")
    add_body(d,
        "Artificial Intelligence (AI) and Large Language Models (LLMs) are changing the way "
        "people use web applications. Many users find it difficult to write SQL queries and "
        "reply to emails manually. These tasks take time and require technical knowledge."
    )
    add_body(d,
        f"The {TITLE} is developed to solve these problems. It allows users to interact with "
        "a MySQL database using simple English instead of writing SQL queries. The system "
        "understands the user’s request, converts it into an SQL query, and shows the required result."
    )
    add_body(d,
        "The project also includes a Gmail Auto-Reply Bot connected with Telegram. It checks "
        "unread emails, creates AI-generated reply drafts, and lets the user review and approve "
        "the reply before sending it. This helps save time while keeping the user in control."
    )
    add_body(d,
        f"This project is developed using {STACK}, Google Gemini AI, Telegram Bot API, and "
        "Gmail API. It provides a simple, fast, and user-friendly solution for database "
        "management and email automation."
    )
    add_h2(d, "Technology Used:")
    for item in [
        "Frontend: React.js (Vite), HTML, CSS, JavaScript",
        "Backend: Node.js and Express.js",
        "Database: MySQL",
        "AI Technology: Google Gemini AI",
        "APIs: Gmail API, Telegram Bot API",
    ]:
        add_bullet(d, item)
    d.add_page_break()

    # Modules
    add_heading_center(d, "3. Project Modules")
    mods = [
        ("1. User Authentication Module",
         "This module allows users to register and log in securely. Only authorized users can access the system and use its features."),
        ("2. AI SQL Agent Module",
         "This module allows users to interact with the database using simple English. The AI converts the user’s request into an SQL query, executes it, and displays the result."),
        ("3. Database Management Module",
         "This module manages all database operations such as adding, updating, deleting, and viewing records in the MySQL database."),
        ("4. Gmail Auto-Reply Module",
         "This module checks unread Gmail messages and generates reply drafts using Google Gemini AI."),
        ("5. Telegram Bot Module",
         "This module sends unread email details and AI-generated replies to Telegram. The user can review the reply and choose to send or skip it."),
        ("6. PDF Report Module",
         "This module generates PDF reports of the database records, allowing users to download and save the data easily."),
    ]
    for h, t in mods:
        add_h2(d, h)
        add_body(d, t)
    d.add_page_break()

    # Gantt
    add_heading_center(d, "4. Project Plan (Gantt Chart)")
    add_body(d, "The project follows a structured development plan across major phases:")
    gantt_rows = [
        ("Sr.", "Task Name", "Weeks"),
        ("1", "Requirement Gathering", "Week 1 – 2"),
        ("2", "Planning", "Week 1 – 4"),
        ("3", "Designing", "Week 3 – 6"),
        ("4", "Coding / Development", "Week 5 – 10"),
        ("5", "Testing & Deployment", "Week 9 – 12"),
    ]
    gt = d.add_table(rows=len(gantt_rows), cols=3)
    gt.style = "Table Grid"
    for i, row in enumerate(gantt_rows):
        for j, val in enumerate(row):
            gt.rows[i].cells[j].text = val
    d.add_page_break()

    # Requirements
    add_heading_center(d, "5. Project Requirements")
    add_h2(d, "Hardware Requirements:")
    for item in [
        "Processor: Intel Core i5 (or equivalent)",
        "RAM: 8 GB minimum (16 GB recommended)",
        "Storage: 256 GB SSD (or higher)",
        "Input Device: Keyboard, Mouse",
        "Output Device: Monitor / LCD / LED, Printer",
        "Internet Connection: Required for Gemini AI, Gmail API, and Telegram Bot",
    ]:
        add_bullet(d, item)
    add_h2(d, "Software Requirements:")
    for item in [
        "Frontend: React.js, HTML, CSS, JavaScript, Vite, Visual Studio Code",
        "Programming Language: JavaScript (Node.js)",
        "Framework: Express.js, React.js",
        "Database: MySQL",
        "AI Technology: Google Gemini AI",
        "API: Gmail API, Telegram Bot API",
        "Platform: Web-Based Application",
    ]:
        add_bullet(d, item)
    d.add_page_break()

    # ER
    add_heading_center(d, "6. E-R Diagram")
    add_body(d,
        "The system uses relational entities centered on the authenticated USER. "
        "Database records are managed through the AI SQL Agent. Email and AI reply entities "
        "support the Gmail–Telegram approval flow."
    )
    add_body(d,
        "Entities: USER → DB_RECORD / EMAIL → PDF_REPORT / AI_REPLY → TELEGRAM_ACTION "
        "(with primary keys and foreign keys as shown in the PDF diagram)."
    )
    d.add_page_break()

    # DFD
    add_heading_center(d, "7. Data Flow Diagram (DFD)")
    add_body(d,
        "User sends a natural language query or email review request to the system "
        "(React frontend + Node/Express backend). The system connects to MySQL, Gemini AI, "
        "Gmail API, and the PDF Report module. AI-generated email replies go to the Telegram "
        "Bot for user approval or rejection, then back to the user."
    )
    d.add_page_break()

    # Conclusion
    add_heading_center(d, "8. Conclusion")
    add_body(d,
        f"The {TITLE} is a smart and user-friendly system that makes database management and "
        "email handling easier. It allows users to interact with a MySQL database using simple "
        "English without writing SQL queries. The system also generates AI-powered email replies "
        "and lets users review them on Telegram before sending."
    )
    add_body(d,
        f"This project saves time, reduces manual work, and improves accuracy. It combines "
        f"{STACK}, Google Gemini AI, Gmail API, and Telegram Bot API to provide an efficient "
        "and secure solution. The project can be further improved by adding more AI features "
        "and support for other messaging platforms in the future."
    )
    d.add_page_break()

    # Bibliography
    add_heading_center(d, "9. Bibliography")
    add_h2(d, "Books")
    add_body(d, "1. Systems Analysis and Design Methods — Jeffrey L. Whitten & Lonnie D. Bentley — McGraw-Hill Education")
    add_body(d, "2. Learning React — Alex Banks & Eve Porcello — O’Reilly Media")
    add_body(d, "3. MySQL Developer’s Library — Paul DuBois — Addison-Wesley Professional")
    add_h2(d, "Websites")
    for w in [
        "React Documentation – https://react.dev",
        "Node.js Documentation – https://nodejs.org",
        "Express.js Documentation – https://expressjs.com",
        "Google Gemini AI Documentation – https://ai.google.dev",
        "MySQL Documentation – https://dev.mysql.com",
        "Gmail API Documentation – https://developers.google.com/gmail/api",
        "Telegram Bot API Documentation – https://core.telegram.org/bots/api",
    ]:
        add_bullet(d, w)

    d.save(DOCX_OUT)
    d.save(DOCX_OUT2)
    print("DOCX:", DOCX_OUT)
    print("DOCX:", DOCX_OUT2)


if __name__ == "__main__":
    build_pdf()
    build_docx()
    # verify no MERN
    with open(PDF_OUT, "rb") as f:
        raw = f.read()
    assert b"MERN" not in raw and b"Mern" not in raw, "MERN still in PDF!"
    print("OK: no MERN in PDF binary")
